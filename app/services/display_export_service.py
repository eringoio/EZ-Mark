"""Export the same localized table markup used by the results view.

Templates are the presentation definition. Explicit numeric column metadata keeps
identifiers and user labels as text while preserving Excel score/percentage types.
"""

import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


@dataclass
class DisplayTable:
    key: str
    title: str
    headers: list[str] = field(default_factory=list)
    formats: list[str | None] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)


class ResultTableParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tables: list[DisplayTable] = []
        self.table = None
        self.cell = None
        self.parts = []
        self.row = []
        self.attrs = {}
        self.header = False

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == "table" and "data-export-key" in attrs:
            self.table = DisplayTable(attrs["data-export-key"], attrs["data-export-title"])
            self.tables.append(self.table)
        if not self.table:
            return
        if tag == "tr":
            self.row = []
            self.header = False
        if tag in {"td", "th"}:
            self.cell = tag
            self.parts = []
            self.attrs = attrs
            self.header |= tag == "th"
        elif self.cell:
            self.parts.append(" ")

    def handle_data(self, data):
        if self.table and self.cell:
            self.parts.append(data)

    def handle_endtag(self, tag):
        if not self.table:
            return
        if tag in {"td", "th"} and self.cell:
            value = " ".join("".join(self.parts).split())
            self.row.append(value)
            if tag == "th":
                self.table.formats.append(self.attrs.get("data-export-number"))
            else:
                self.row.extend([""] * (int(self.attrs.get("colspan", 1)) - 1))
            self.cell = None
        elif tag == "tr":
            if self.header:
                self.table.headers = self.row
            elif self.row:
                self.table.rows.append(self.row)
        elif tag == "table":
            self.table = None


def build_display_workbook(tables: list[DisplayTable]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    used = set()
    for table in tables:
        base = re.sub(r"[\\/*?:\[\]\x00-\x1f]", "_", table.title).strip("'") or "Results"
        title = base[:31]
        suffix = 1
        while title.casefold() in used:
            suffix += 1
            ending = f" ({suffix})"
            title = base[: 31 - len(ending)] + ending
        used.add(title.casefold())
        sheet = workbook.create_sheet(title)
        columns = [index for index, heading in enumerate(table.headers) if heading]
        for out_col, index in enumerate(columns, 1):
            cell = sheet.cell(1, out_col, table.headers[index])
            cell.data_type = "s"
            cell.font = Font(bold=True)
        for out_row, row in enumerate(table.rows, 2):
            for out_col, index in enumerate(columns, 1):
                value = row[index] if index < len(row) else ""
                fmt = table.formats[index]
                cell = sheet.cell(out_row, out_col)
                if value in {"", "—"}:
                    continue
                if fmt:
                    # Score fractions retain a numeric numerator with a visible denominator.
                    match = re.fullmatch(r"(-?\d+(?:\.\d+)?)(%|\s*/\s*\d+(?:\.\d+)?)?", value)
                    if match:
                        number = float(match[1])
                        if match[2] == "%":
                            number /= 100
                        elif match[2]:
                            fmt = fmt + '"' + match[2] + '"'
                        cell.value = number
                        cell.number_format = fmt
                        continue
                cell.value = value
                cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            width = max(
                (
                    len(str(cell.value or "")) + sum(ord(c) > 127 for c in str(cell.value or ""))
                    for cell in column
                ),
                default=10,
            )
            sheet.column_dimensions[get_column_letter(column[0].column)].width = min(
                64, max(14, width + 2)
            )
    if not workbook.worksheets:
        workbook.create_sheet("Results")
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
