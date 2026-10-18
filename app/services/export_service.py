import csv
from io import StringIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import get_settings
from app.database import data_relative_path, ensure_data_directories
from app.models import Export, MockExam, Question, Response


def generate_student_summary_export(db: Session, exam: MockExam, file_type: str) -> str:
    rows = [_student_summary_row(exam, result) for result in exam.results]
    headers = _student_summary_headers(exam)
    return _write_export(db, exam, "student_summary", file_type, headers, rows)


def generate_per_question_export(db: Session, exam: MockExam, file_type: str) -> str:
    responses = list(
        db.scalars(select(Response).join(Response.question).where(Question.exam_id == exam.id))
    )
    rows = [_per_question_row(response, exam) for response in responses]
    headers = [
        "exam_name",
        "student_number",
        "name",
        "class_name",
        "physical_sheet_no",
        "physical_sheet_name",
        "question_no",
        "display_label",
        "test_section",
        "scoring_section",
        "item_type",
        "choice_count",
        "correct_answer",
        "detected_answer",
        "final_answer",
        "is_correct",
        "confidence",
        "response_status",
        "reviewed_manually",
    ]
    return _write_export(db, exam, "per_question", file_type, headers, rows)


def generate_question_analysis_export(db: Session, exam: MockExam, file_type: str) -> str:
    from app.services.scoring_service import question_analysis

    rows = []
    for row in question_analysis(db, exam):
        question = row["question"]
        rows.append(
            {
                "exam_name": exam.name,
                "question_no": question.question_no,
                "display_label": question.display_label,
                "physical_sheet_name": question.physical_sheet.sheet_name,
                "test_section": question.test_section,
                "scoring_section": question.scoring_section.name
                if question.scoring_section
                else "",
                "item_type": question.item_type,
                "correct_answer": question.correct_answer,
                "student_count": row["student_count"],
                "answered_count": row["answered_count"],
                "correct_count": row["correct_count"],
                "correct_rate": row["correct_rate"],
                "blank_count": row["blank_count"],
                "multiple_mark_count": row["multiple_mark_count"],
            }
        )
    headers = [
        "exam_name",
        "question_no",
        "display_label",
        "physical_sheet_name",
        "test_section",
        "scoring_section",
        "item_type",
        "correct_answer",
        "student_count",
        "answered_count",
        "correct_count",
        "correct_rate",
        "blank_count",
        "multiple_mark_count",
    ]
    return _write_export(db, exam, "question_analysis", file_type, headers, rows)


def export_path(relative_path: str) -> Path:
    path = get_settings().data_root / relative_path
    if not path.exists() or not path.is_file():
        raise FileNotFoundError(relative_path)
    return path


def _student_summary_headers(exam: MockExam) -> list[str]:
    headers = ["exam_name", "exam_date", "student_number", "name", "class_name"]
    for section in exam.scoring_sections:
        headers.extend(
            [
                f"{section.code}_raw_correct",
                f"{section.code}_raw_total",
                f"{section.code}_accuracy",
                f"{section.code}_score",
            ]
        )
    if exam.pass_mark_value is not None:
        headers.extend(["pass_mark_value", "pass_mark_type"])
    headers.extend(["total_score", "pass_fail_status"])
    return headers


def _student_summary_row(exam: MockExam, result) -> dict:
    row = {
        "exam_name": exam.name,
        "exam_date": exam.exam_date,
        "student_number": result.student.student_number,
        "name": result.student.name,
        "class_name": result.student.group_for_exam(exam.id),
        "total_score": result.total_score,
        "pass_fail_status": result.pass_fail_status,
    }
    if exam.pass_mark_value is not None:
        row["pass_mark_value"] = exam.pass_mark_value
        row["pass_mark_type"] = exam.pass_mark_type
    for section in exam.scoring_sections:
        raw = result.raw_json.get(section.code, {})
        row[f"{section.code}_raw_correct"] = raw.get("correct")
        row[f"{section.code}_raw_total"] = raw.get("total")
        row[f"{section.code}_accuracy"] = raw.get("accuracy")
        score_details = result.score_details_json.get(section.code)
        # Legacy scalar payloads store the section score directly; current payloads
        # store {earned, possible, percentage}.
        row[f"{section.code}_score"] = (
            score_details.get("earned") if isinstance(score_details, dict) else score_details
        )
    return row
    return row


def _per_question_row(response: Response, exam: MockExam) -> dict:
    page = response.scanned_page
    instance = page.sheet_instance
    student = instance.student if instance else None
    question = response.question
    return {
        "exam_name": exam.name,
        "student_number": student.student_number if student else "",
        "name": student.name if student else "",
        "class_name": student.group_for_exam(exam.id) if student else "",
        "physical_sheet_no": question.physical_sheet.sheet_no,
        "physical_sheet_name": question.physical_sheet.sheet_name,
        "question_no": question.question_no,
        "display_label": question.display_label,
        "test_section": question.test_section,
        "scoring_section": question.scoring_section.name if question.scoring_section else "",
        "item_type": question.item_type,
        "choice_count": question.choice_count,
        "correct_answer": question.correct_answer,
        "detected_answer": response.detected_answer,
        "final_answer": response.final_answer,
        "is_correct": response.is_correct,
        "confidence": response.confidence,
        "response_status": response.status,
        "reviewed_manually": bool(response.reviewed_at),
    }


def _write_export(
    db: Session,
    exam: MockExam,
    export_type: str,
    file_type: str,
    headers: list[str],
    rows: list[dict],
) -> str:
    if file_type not in {"csv", "xlsx"}:
        raise ValueError("Export type must be csv or xlsx.")
    ensure_data_directories()
    output_dir = get_settings().data_root / "exports" / f"exam_{exam.id}"
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"{_safe_filename(exam.name)}_{export_type}.{file_type}"
    if file_type == "csv":
        _write_csv(path, headers, rows)
    else:
        _write_xlsx(path, headers, rows)
    relative_path = data_relative_path(path)
    db.add(Export(exam_id=exam.id, export_type=export_type, file_path=relative_path))
    db.commit()
    return relative_path


def _write_csv(path: Path, headers: list[str], rows: list[dict]) -> None:
    buffer = StringIO()
    writer = csv.DictWriter(buffer, fieldnames=headers, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(rows)
    path.write_text("\ufeff" + buffer.getvalue(), encoding="utf-8")


def _write_xlsx(path: Path, headers: list[str], rows: list[dict]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Export"
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header) for header in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = cell.font.copy(bold=True)
    for column_index, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(column_index)].width = max(12, len(header) + 2)
    workbook.save(path)


def _safe_filename(value: str) -> str:
    safe = "".join(char if char.isalnum() else "_" for char in value.strip())
    return safe or "assessment"
