import csv
import json
from dataclasses import asdict, dataclass
from io import BytesIO, StringIO
from typing import Any
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import ExamStudent, MockExam, Student
from app.services.i18n_service import _, i18n

REQUIRED_STUDENT_HEADERS = ("student_number", "name", "class_name")
# Public file headings are separate from the internal model/payload field names.
HEADER_ALIASES = {
    "student number": "student_number",
    "学生番号": "student_number",
    "name": "name",
    "氏名": "name",
    "group": "class_name",
    "グループ": "class_name",
}


@dataclass(frozen=True)
class StudentImportRow:
    row_number: int
    student_number: str
    name: str
    class_name: str
    status: str
    message: str

    existing_id: int | None = None
    existing_name: str | None = None
    existing_group: str | None = None
    link_state: tuple | None = None
    requested_seat_order: int | None = None

    @property
    def is_valid(self) -> bool:
        return self.status in {"ok", "already_exists", "conflict", "attached"}


@dataclass(frozen=True)
class StudentImportPreview:
    rows: list[StudentImportRow]
    missing_headers: list[str]

    @property
    def valid_rows(self) -> list[StudentImportRow]:
        return [row for row in self.rows if row.is_valid]

    @property
    def error_rows(self) -> list[StudentImportRow]:
        return [row for row in self.rows if not row.is_valid]


def parse_student_import(
    filename: str, content: bytes, db: Session, exam: MockExam | None = None
) -> StudentImportPreview:
    extension = filename.lower().rsplit(".", 1)[-1]
    try:
        if extension == "csv":
            raw_rows = _read_csv_rows(content)
        elif extension == "xlsx":
            raw_rows = _read_xlsx_rows(content)
        else:
            raise ValueError(_("students.conflicts.file_type"))
    except (BadZipFile, InvalidFileException, UnicodeDecodeError) as exc:
        raise ValueError(_("students.conflicts.invalid_file")) from exc
    return validate_student_rows(raw_rows, db, exam)


def validate_student_rows(
    raw_rows: list[dict[str, Any]], db: Session, exam: MockExam | None = None
) -> StudentImportPreview:
    if not raw_rows:
        return StudentImportPreview(rows=[], missing_headers=list(REQUIRED_STUDENT_HEADERS))
    raw_rows = [_normalize_student_row(row) for row in raw_rows]
    headers = set(raw_rows[0].keys())
    missing_headers = [header for header in REQUIRED_STUDENT_HEADERS if header not in headers]
    if missing_headers:
        return StudentImportPreview(rows=[], missing_headers=missing_headers)

    existing_students: dict[str, list[Student]] = {}
    for student in db.scalars(select(Student)):
        existing_students.setdefault("".join(student.student_number.split()), []).append(student)
    seen_numbers: set[str] = set()
    rows: list[StudentImportRow] = []
    for row_number, raw in enumerate(raw_rows, start=2):
        student_number = str(raw.get("student_number") or "").strip()
        name = str(raw.get("name") or "").strip()
        class_name = str(raw.get("class_name") or "").strip()
        matches = existing_students.get("".join(student_number.split()), [])
        existing = matches[0] if len(matches) == 1 else None
        links = (
            [link for link in existing.exam_links if exam and link.exam_id == exam.id]
            if existing
            else []
        )
        link = links[0] if len(links) == 1 else None
        group = link.group_name if link else (existing.class_name if existing else None)
        key = "".join(student_number.split())
        status = "ok"
        message = _("students.conflicts.ok")
        if not student_number:
            status = "error"
            message = _("students.conflicts.missing_number")
        elif not name:
            status = "error"
            message = _("students.conflicts.missing_name")
        elif not class_name:
            status = "error"
            message = _("students.conflicts.missing_group")
        elif key in seen_numbers:
            status = "error"
            message = _("students.conflicts.duplicate")
        elif len(matches) > 1 or len(links) > 1:
            status = "error"
            message = _("students.conflicts.ambiguous")
        elif existing:
            status = (
                "conflict" if existing.name != name or group != class_name else "already_exists"
            )
            if status != "conflict" and link and link.included:
                status = "attached"
            message = _("students.conflicts." + status)
        seen_numbers.add(key)
        rows.append(
            StudentImportRow(
                row_number=row_number,
                student_number=student_number,
                name=name,
                class_name=class_name,
                status=status,
                message=message,
                requested_seat_order=raw.get("requested_seat_order"),
                existing_id=existing.id if existing else None,
                existing_name=existing.name if existing else None,
                existing_group=group,
                link_state=(link.id, link.included, link.seat_order, link.group_name)
                if link
                else None,
            )
        )
    return StudentImportPreview(rows=rows, missing_headers=[])


def import_preview_to_payload(preview: StudentImportPreview) -> str:
    return json.dumps([asdict(row) for row in preview.valid_rows], ensure_ascii=False)


def import_students_from_payload(
    db: Session, exam: MockExam, payload: str, choices: dict | None = None, *, commit: bool = True
) -> int:
    from app.services.exam_service import get_run_or_404

    get_run_or_404(db, exam.id, for_write=True)
    rows = json.loads(payload)
    if not isinstance(rows, list) or not rows:
        raise ValueError(_("students.conflicts.invalid"))
    choices = choices or {}
    allowed_rows = {str(row["row_number"]) for row in rows}
    if set(choices) - allowed_rows:
        raise ValueError(_("students.conflicts.invalid"))
    # Validate every row and snapshot before making any changes.
    fresh = validate_student_rows(rows, db, exam)
    if fresh.missing_headers or fresh.error_rows or len(fresh.rows) != len(rows):
        raise ValueError(_("students.conflicts.invalid"))
    for old, new in zip(rows, fresh.rows, strict=True):
        for key in ("existing_id", "existing_name", "existing_group", "link_state"):
            expected = old.get(key)
            actual = getattr(new, key)
            if key == "link_state" and actual:
                actual = list(actual)
            if expected != actual:
                raise ValueError(_("students.conflicts.stale"))
        decision = choices.get(str(old["row_number"]), {})
        if (
            set(decision) - {"name", "group", "skip"}
            or decision.get("name", "keep") not in {"keep", "update"}
            or decision.get("group", "keep") not in {"keep", "import"}
            or decision.get("skip", "no") not in {"yes", "no"}
        ):
            raise ValueError(_("students.conflicts.invalid"))
    imported = 0
    next_order = max((link.seat_order or 0 for link in exam.students), default=0) + 1
    try:
        for row in rows:
            decision = choices.get(str(row["row_number"]), {})
            if decision.get("skip") == "yes":
                continue
            student = db.get(Student, row["existing_id"]) if row.get("existing_id") else None
            if student is None:
                student = Student(
                    student_number=row["student_number"],
                    name=row["name"],
                    class_name=row["class_name"],
                )
                db.add(student)
                db.flush()
            elif decision.get("name") == "update":
                student.name = row["name"]
            link = db.scalar(
                select(ExamStudent).where(
                    ExamStudent.exam_id == exam.id, ExamStudent.student_id == student.id
                )
            )
            group = (
                row["class_name"]
                if decision.get("group") == "import" or not row.get("existing_id")
                else row["existing_group"]
            )
            if link is None:
                db.add(
                    ExamStudent(
                        exam_id=exam.id,
                        student_id=student.id,
                        class_name=group,
                        seat_order=row.get("requested_seat_order")
                        if row.get("requested_seat_order") is not None
                        else next_order,
                        included=True,
                    )
                )
                next_order += 1
                imported += 1
            else:
                imported += int(not link.included)
                link.included = True
                link.class_name = group
                if row.get("requested_seat_order") is not None:
                    link.seat_order = row["requested_seat_order"]
        db.flush()
        if commit:
            db.commit()
    except Exception:
        db.rollback()
        raise
    return imported


def build_student_template_xlsx(locale: str | None = None) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = i18n.translate("exams.detail.students", locale)
    sheet.append(
        [
            i18n.translate("students.table." + key, locale)
            for key in ("student_number", "name", "class")
        ]
    )
    sheet.append(["SYN001", "Student A", "A"])
    sheet.append(["SYN002", "Student B", "A"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for column in ("A", "B", "C"):
        sheet.column_dimensions[column].width = 24
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _validate_headers(headers: list[str]) -> list[str]:
    unsupported = [header for header in headers if header.strip().lower() not in HEADER_ALIASES]
    if unsupported:
        raise ValueError(
            _("students.conflicts.unsupported_headers", headers=", ".join(unsupported))
        )
    normalized = [_normalize_header(header) for header in headers]
    if len(normalized) != len(set(normalized)):
        raise ValueError(_("students.conflicts.duplicate_headers"))
    return normalized


def _read_csv_rows(content: bytes) -> list[dict[str, Any]]:
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(StringIO(text))
    if reader.fieldnames is None:
        return []
    reader.fieldnames = _validate_headers(reader.fieldnames)
    rows = list(reader)
    if any(None in row or any(value is None for value in row.values()) for row in rows):
        raise ValueError(_("students.conflicts.invalid_rows"))
    return rows


def _read_xlsx_rows(content: bytes) -> list[dict[str, Any]]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    if not rows:
        return []
    headers = _validate_headers([str(cell or "") for cell in rows[0]])
    parsed_rows = []
    for row in rows[1:]:
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        parsed_rows.append(
            {
                header: "" if value is None else value
                for header, value in zip(headers, row, strict=False)
                if header
            }
        )
    return parsed_rows


def _normalize_header(header: str) -> str:
    key = (header or "").strip().lower()
    return HEADER_ALIASES.get(key, key)


def _normalize_student_row(row: dict[str, Any]) -> dict[str, Any]:
    normalized = {_normalize_header(key): value for key, value in row.items()}
    if len(normalized) != len(row):
        raise ValueError(_("students.conflicts.duplicate_headers"))
    return normalized
