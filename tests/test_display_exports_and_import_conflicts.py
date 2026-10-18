from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, select, text
from sqlalchemy.orm import Session

from app.database import PUBLIC_SCHEMA_VERSION, Base, SessionLocal
from app.main import create_app
from app.models import ExamStudent, MockExam, Student, StudentResult
from app.services.display_export_service import DisplayTable, build_display_workbook
from app.services.i18n_service import i18n
from app.services.student_import_service import (
    build_student_template_xlsx,
    import_preview_to_payload,
    import_students_from_payload,
    parse_student_import,
)


@pytest.fixture
def db():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with engine.begin() as connection:
        connection.execute(
            text("INSERT INTO schema_info (schema_version) VALUES (:v)"),
            {"v": PUBLIC_SCHEMA_VERSION},
        )
    with Session(engine) as session:
        yield session
    engine.dispose()


@pytest.mark.parametrize(
    "locale,headers",
    [
        ("en", ["Student Number", "Name", "Group"]),
        ("ja", ["学生番号", "氏名", "グループ"]),
    ],
)
def test_localized_three_column_template_roundtrip(db, locale, headers):
    content = build_student_template_xlsx(locale)
    sheet = load_workbook(BytesIO(content)).active
    assert [cell.value for cell in sheet[1]] == headers
    assert sheet.max_column == 3
    preview = parse_student_import("test.xlsx", content, db)
    assert len(preview.valid_rows) == 2
    assert not preview.error_rows


@pytest.mark.parametrize(
    "headers",
    ["Student Number,Name,Group", "学生番号,氏名,グループ"],
)
def test_header_aliases_in_any_locale(db, headers):
    preview = parse_student_import("test.csv", (headers + "\nSYN1,Student A,A\n").encode(), db)
    assert preview.valid_rows[0].class_name == "A"


@pytest.mark.parametrize(
    "headers", ["Student Number,Student Code,Name,Group", "student_number,name,group,unknown"]
)
def test_duplicate_ambiguous_and_unknown_headers(db, headers):
    with pytest.raises(ValueError):
        parse_student_import("test.csv", (headers + "\nSYN1,SYN2,Student A,A\n").encode(), db)


def _conflict(db):
    student = Student(student_number="SYN1", name="Student A", class_name="A")
    earlier = MockExam(name="Earlier")
    current = MockExam(name="Current")
    db.add_all([student, earlier, current])
    db.flush()
    db.add(ExamStudent(exam_id=earlier.id, student_id=student.id, seat_order=77))
    db.add(StudentResult(exam_id=earlier.id, student_id=student.id, total_score=42))
    db.commit()
    preview = parse_student_import(
        "test.csv", b"Student Number,Name,Group\nSYN1,Student Revised,B\n", db, current
    )
    assert preview.rows[0].status == "conflict"
    return student, earlier, current, import_preview_to_payload(preview)


@pytest.mark.parametrize(
    "decision,name,group,count",
    [
        ({}, "Student A", "A", 1),
        ({"name": "update", "group": "import"}, "Student Revised", "B", 1),
        ({"skip": "yes"}, "Student A", "A", 0),
    ],
)
def test_conflict_decisions_preserve_history(db, decision, name, group, count):
    student, earlier, current, payload = _conflict(db)
    assert import_students_from_payload(db, current, payload, {"2": decision}) == count
    assert student.name == name
    assert student.class_name == "A"
    assert student.group_for_exam(earlier.id) == "A"
    assert earlier.results[0].total_score == 42
    if count:
        assert student.group_for_exam(current.id) == group


def test_stale_and_invalid_decisions_cannot_partially_import(db):
    student, earlier, current, payload = _conflict(db)
    with pytest.raises(ValueError):
        import_students_from_payload(db, current, payload, {"2": {"name": "unexpected"}})
    student.name = "Concurrent correction"
    db.commit()
    with pytest.raises(ValueError):
        import_students_from_payload(db, current, payload)
    assert not current.students


def test_duplicate_database_identifiers_fail_loudly(db):
    for name in ("Student A", "Student B"):
        db.add(Student(student_number="SYN1", name=name, class_name="A"))
    db.commit()
    preview = parse_student_import("test.csv", b"Student Number,Name,Group\nSYN1,Student A,A\n", db)
    assert not preview.valid_rows
    assert preview.error_rows


def test_reincluded_student_keeps_custom_order(db):
    student, earlier, current, _ = _conflict(db)
    link = earlier.students[0]
    link.included = False
    db.commit()
    preview = parse_student_import(
        "test.csv", b"Student Number,Name,Group\nSYN1,Student A,A\n", db, earlier
    )
    import_students_from_payload(db, earlier, import_preview_to_payload(preview))
    assert link.included
    assert link.seat_order == 77
    assert len(earlier.students) == 1


def test_workbook_safety_precision_missing_and_sheet_names():
    table = DisplayTable(
        "test",
        "日本語/" + "長" * 40,
        ["学生番号", "得点", "割合", "氏名"],
        [None, "0.00", "0%", None],
        [
            ["00123", "4.25", "75%", "=1+1"],
            ["002", "—", "—", "@SUM(A1)"],
        ],
    )
    workbook = load_workbook(BytesIO(build_display_workbook([table, table])))
    assert len(workbook.sheetnames) == 2
    assert len(set(workbook.sheetnames)) == 2
    sheet = workbook.worksheets[0]
    assert sheet["A2"].value == "00123"
    assert sheet["B2"].value == 4.25
    assert sheet["C2"].value == 0.75
    assert sheet["C2"].number_format == "0%"
    assert sheet["D2"].data_type == "s"
    assert sheet["B3"].value is None
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref


@pytest.mark.parametrize("locale", ["en", "ja"])
def test_display_export_routes_match_localized_screen(locale):
    with TestClient(create_app()) as client:
        with SessionLocal() as session:
            exam = MockExam(
                name="Synthetic Export",
                assessment_type="generic_test",
                pass_mark_value=50,
                pass_mark_type="percentage",
            )
            student = Student(
                student_number="EXPORT-" + locale, name="=Synthetic", class_name="試験組"
            )
            session.add_all([exam, student])
            session.flush()
            session.add(ExamStudent(exam_id=exam.id, student_id=student.id, class_name="今回組"))
            session.add(
                StudentResult(
                    exam_id=exam.id,
                    student_id=student.id,
                    total_score=75,
                    score_details_json={"max_points": 100, "percentage": 75},
                    pass_fail_status="pass",
                )
            )
            session.commit()
            exam_id = exam.id
        screen = client.get(f"/exams/{exam_id}/results?lang={locale}")
        response = client.get(f"/exams/{exam_id}/exports/display/all.xlsx?lang={locale}")
        assert response.status_code == 200
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.worksheets[0]
        assert sheet["B1"].value == i18n.translate("students.table.student_number", locale)
        assert sheet["B1"].value in screen.text
        assert sheet["D2"].value == "今回組"
        assert sheet["E2"].value == 75
        assert sheet["F2"].value == 0.75
        assert sheet["G2"].value == i18n.translate("results.status.pass", locale)
        assert sheet["C2"].data_type == "s"
        assert (
            client.get(f"/exams/{exam_id}/students/import-template.xlsx?lang={locale}").status_code
            == 200
        )


def test_confirmation_is_server_bound_atomic_and_idempotent():
    import re

    with TestClient(create_app()) as client:
        with SessionLocal() as session:
            student = Student(student_number="CONFIRM-SYN1", name="Student Before", class_name="A")
            exam = MockExam(name="Confirm Synthetic")
            other = MockExam(name="Other Synthetic")
            session.add_all([student, exam, other])
            session.commit()
            exam_id, other_id, student_id = exam.id, other.id, student.id
        data = {"student_number": "CONFIRM-SYN1", "name": "Student After", "class_name": "B"}
        preview = client.post(f"/exams/{exam_id}/students/add-manual", data=data)
        assert preview.status_code == 200
        token = re.search(r'name="import_payload" value="([^"]+)"', preview.text)[1]
        assert "Student Before" in preview.text and "Student After" in preview.text
        assert (
            client.post(
                f"/exams/{other_id}/students/import/confirm", data={"import_payload": token}
            ).status_code
            == 400
        )
        assert (
            client.post(
                f"/exams/{exam_id}/students/import/confirm",
                data={"import_payload": token, "name_2": "bad"},
            ).status_code
            == 400
        )
        decision = {"import_payload": token, "name_2": "update", "group_2": "import"}
        assert (
            client.post(
                f"/exams/{exam_id}/students/import/confirm", data=decision, follow_redirects=False
            ).status_code
            == 303
        )
        assert (
            client.post(
                f"/exams/{exam_id}/students/import/confirm", data=decision, follow_redirects=False
            ).status_code
            == 303
        )
        with SessionLocal() as session:
            student = session.get(Student, student_id)
            assert student.name == "Student After"
            assert student.class_name == "A"
            links = list(session.scalars(select(ExamStudent).where(ExamStudent.exam_id == exam_id)))
            assert len(links) == 1
            assert links[0].group_name == "B"
        # Client-provided serialized roster data is never accepted as a preview token.
        assert (
            client.post(
                f"/exams/{exam_id}/students/import/confirm",
                data={"import_payload": '[{"student_number":"forged"}]'},
            ).status_code
            == 400
        )


def test_peer_review_display_exports_preserve_names_and_numeric_ratings():
    from tests.test_peer_review_results_page import _peer_run_with_ratings

    with TestClient(create_app()) as client:
        exam_id = _peer_run_with_ratings(client)
        with SessionLocal() as session:
            from app.models import PeerReviewForm

            form = session.scalar(select(PeerReviewForm).where(PeerReviewForm.exam_id == exam_id))
            group = form.rater_groups[0]
            group.name = "評価組/日本語"
            competency = form.exam.competencies[0]
            competency.name = "日本語能力"
            session.commit()
            from app.services.peer_review_scoring_service import (
                recalculate_anonymous_peer_review_results,
            )

            recalculate_anonymous_peer_review_results(session, form)
        response = client.get(f"/exams/{exam_id}/exports/display/all.xlsx?lang=ja")
        assert response.status_code == 200
        workbook = load_workbook(BytesIO(response.content))
        assert len(workbook.worksheets) >= 2
        group_sheet = next(sheet for sheet in workbook if sheet.title.startswith("評価組"))
        assert "日本語能力" in [cell.value for cell in group_sheet[1]]
        assert group_sheet["C2"].value == 4.5
        assert group_sheet["D2"].value == 2


def test_roster_rows_survive_repeated_init_and_null_group_falls_back(monkeypatch):
    """Public v1 databases are opened as-is across restarts, with the roster
    group fallback (``ExamStudent.group_name``) kept stable for pre-snapshot
    rows whose ``class_name`` is NULL."""
    from sqlalchemy import select, text

    from app import database
    from app.models import ExamStudent, Student

    engine = create_engine("sqlite:///:memory:")
    monkeypatch.setattr(database, "engine", engine)
    database.init_database()
    with engine.begin() as connection:
        connection.execute(
            text(
                "INSERT INTO students (student_number, name, class_name, active, "
                "created_at, updated_at) VALUES ('R00', 'Historic A', 'A', 1, '2026-01-01 00:00:00', '2026-01-01 00:00:00')"
            )
        )
        connection.execute(
            text(
                "INSERT INTO exam_students (exam_id, student_id, class_name, included, "
                "created_at) VALUES (3, 1, NULL, 1, '2026-01-01 00:00:00')"
            )
        )
    database.init_database()
    database.init_database()
    from sqlalchemy.orm import Session

    with Session(engine) as session:
        link = session.scalar(select(ExamStudent).where(ExamStudent.student_id == 1))
        assert (link.exam_id, link.class_name) == (3, None)
        student = session.get(Student, 1)
        assert student.class_name == "A"
        assert link.group_name == "A"
        # A later class change must not rewrite the stored per-roster snapshot,
        # and the property fallback keeps the current class visible.
        student.class_name = "Default B"
        session.commit()
        link_after = session.scalar(select(ExamStudent).where(ExamStudent.student_id == 1))
        assert link_after.class_name is None
        assert link_after.group_name == "Default B"
    engine.dispose()


# (section_code, item_type, correct_answer) — a small fully-scored generic exam.
_QUESTIONS = [
    ("language_knowledge", "moji", "1"),
    ("language_knowledge", "goi", "2"),
    ("listening", "listening", "3"),
]


def _build_scored_exam(session, qr_token: str = "tok-ss1"):
    from app.models import (
        ExamStudent,
        MockExam,
        PhysicalSheet,
        Question,
        Response,
        ScanBatch,
        ScannedPage,
        ScoringSection,
        SheetInstance,
        SheetTemplate,
        Student,
    )
    from app.services.scoring_service import recalculate_exam_results

    exam = MockExam(name="Score Sheet Test", assessment_type="generic_test")
    sheet = PhysicalSheet(sheet_no=1, sheet_name="S", slug="s", display_order=1, required=True)
    lk = ScoringSection(code="language_knowledge", name="LK")
    listening = ScoringSection(code="listening", name="Listening")
    exam.physical_sheets.append(sheet)
    exam.scoring_sections.extend([lk, listening])
    student = Student(student_number="SS1", name="Student SS1", class_name="A")
    session.add_all([exam, student])
    session.commit()
    session.add(ExamStudent(exam_id=exam.id, student_id=student.id))
    template = SheetTemplate(exam_id=exam.id, physical_sheet_id=sheet.id)
    session.add(template)
    session.flush()
    section_by_code = {s.code: s for s in exam.scoring_sections}
    questions = []
    for index, (section_code, item_type, answer) in enumerate(_QUESTIONS, start=1):
        questions.append(
            Question(
                exam_id=exam.id,
                physical_sheet_id=sheet.id,
                scoring_section_id=section_by_code[section_code].id,
                display_label=str(index),
                choice_count=4,
                correct_answer=answer,
                item_type=item_type,
                display_order=index,
            )
        )
    instance = SheetInstance(
        exam_id=exam.id,
        student_id=student.id,
        physical_sheet_id=sheet.id,
        sheet_template_id=template.id,
        qr_token=qr_token,
        printed_name=student.name,
        printed_student_number=student.student_number,
    )
    batch = ScanBatch(exam_id=exam.id, uploaded_file_name="s.pdf", uploaded_file_path="x")
    page = ScannedPage(scan_batch=batch, sheet_instance=instance, page_no=1, image_path="x")
    session.add_all([*questions, instance, batch, page])
    session.flush()
    for question, (_, _, answer) in zip(questions, _QUESTIONS, strict=True):
        session.add(
            Response(scanned_page=page, question=question, final_answer=answer, status="ok")
        )
    session.commit()

    recalculate_exam_results(session, exam)
    return exam


@pytest.mark.parametrize("locale", ["en", "ja"])
def test_all_workbook_includes_localized_analysis_and_priority(locale):
    from tests.test_display_exports_and_import_conflicts import _build_scored_exam

    with TestClient(create_app()) as client:
        with SessionLocal() as session:
            exam = _build_scored_exam(session, qr_token="display-fixture-" + locale)
            exam.students[0].student.student_number = "DISPLAY-" + locale
            session.commit()
            exam_id = exam.id
        workbook = load_workbook(
            BytesIO(client.get(f"/exams/{exam_id}/exports/display/all.xlsx?lang={locale}").content)
        )
        assert len(workbook.worksheets) == 3
        ranking, priorities, analysis = workbook.worksheets
        assert ranking["B1"].value == i18n.translate("students.table.student_number", locale)
        assert isinstance(ranking.cell(2, ranking.max_column - 1).value, (int, float))
        assert priorities["A1"].value == i18n.translate("results.priority.priority", locale)
        assert priorities["A2"].value in {
            i18n.translate("results.priority.band." + key, locale)
            for key in ("high", "medium", "low")
        }
        assert analysis["D2"].data_type == "n"
        assert analysis["D2"].number_format == "0%"
        # Existing technical endpoints continue to use canonical machine headings.
        raw = client.get(f"/exams/{exam_id}/exports/student-summary.xlsx?lang={locale}")
        assert load_workbook(BytesIO(raw.content)).active["A1"].value == "exam_name"


@pytest.mark.parametrize(
    ("key", "file_type"),
    [
        ("student-summary", "csv"),
        ("student-summary", "xlsx"),
        ("per-question", "csv"),
        ("per-question", "xlsx"),
        ("question-analysis", "csv"),
        ("question-analysis", "xlsx"),
    ],
)
def test_technical_export_routes_cover_csv_and_xlsx(key, file_type):
    """Every technical export table is downloadable as plain CSV and XLSX."""
    import csv
    from io import StringIO

    from tests.test_display_exports_and_import_conflicts import _build_scored_exam

    with TestClient(create_app()) as client:
        with SessionLocal() as session:
            exam = _build_scored_exam(session, qr_token=f"export-{key}-{file_type}")
            exam_id = exam.id
        response = client.get(f"/exams/{exam_id}/exports/{key}.{file_type}")
        assert response.status_code == 200, response.text[:300]
        if file_type == "csv":
            assert response.headers["content-type"].startswith("text/csv")
            rows = list(csv.reader(StringIO(response.content.decode("utf-8-sig"))))
            assert rows[0][0] == "exam_name"
            assert len(rows) >= 2, "one data row expected for the scored fixture"
        else:
            workbook = load_workbook(BytesIO(response.content))
            assert workbook.active["A1"].value == "exam_name"


@pytest.mark.parametrize(
    "header",
    [
        "student_number",
        "Student Code",
        "class_name",
        "name_kana",
        "school_management_id",
        "course",
        "nationality",
        "memo",
    ],
)
def test_legacy_headings_and_optional_columns_are_rejected(db, header):
    with pytest.raises(ValueError, match="Unsupported columns"):
        parse_student_import(
            "students.csv",
            f"Student Number,Name,Group,{header}\nSYN99,Student Fiction,A,unused\n".encode(),
            db,
        )


def test_reordered_localized_headers_and_missing_required_column(db):
    preview = parse_student_import(
        "students.csv", "グループ,氏名,学生番号\nA,Student Fiction,SYN99\n".encode(), db
    )
    assert not preview.error_rows
    assert preview.valid_rows[0].student_number == "SYN99"
    missing = parse_student_import("students.csv", b"Name,Group\nStudent Fiction,A\n", db)
    assert missing.missing_headers == ["student_number"]


def test_existing_qr_and_saved_maps_survive_restart(db, monkeypatch):
    """Restarting (re-running init_database) against a public v1 database keeps
    QR tokens, saved coordinate maps and results intact."""
    from copy import deepcopy

    from app import database
    from app.models import SheetInstance
    from tests.test_display_exports_and_import_conflicts import _build_scored_exam

    exam = _build_scored_exam(db)
    sheet = db.scalar(select(SheetInstance).where(SheetInstance.exam_id == exam.id))
    template = sheet.sheet_template
    template.coordinate_map_json = {
        "questions": [{"question_id": exam.questions[0].id, "x_mm": 25, "y_mm": 30}]
    }
    db.commit()
    saved = (sheet.id, sheet.student_id, sheet.qr_token, deepcopy(template.coordinate_map_json))
    result_ids = [result.id for result in exam.results]
    monkeypatch.setattr(database, "engine", db.get_bind())
    database.init_database()
    db.expire_all()
    mapped = db.scalar(select(SheetInstance).where(SheetInstance.qr_token == saved[2]))
    assert (
        mapped.id,
        mapped.student_id,
        mapped.qr_token,
        mapped.sheet_template.coordinate_map_json,
    ) == saved
    assert [result.id for result in exam.results] == result_ids
    assert exam.students[0].group_name == "A"
